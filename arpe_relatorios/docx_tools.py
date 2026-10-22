"""Utilitários para editar o Template.docx no lugar (python-docx + lxml)."""
from __future__ import annotations

import copy
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
VERMELHOS = {"ff0000", "c00000", "e06666"}


def strip_ids(el):
    """Remove w14:paraId/textId das cópias (evita IDs duplicados no Word)."""
    for e in el.iter():
        for a in list(e.attrib):
            if a.startswith("{%s}" % W14):
                del e.attrib[a]
    return el


def clone(el):
    return strip_ids(copy.deepcopy(el))


def texto_el(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


class Editor:
    """Envolve um Document e oferece busca por texto, clonagem e inserção de blocos."""

    def __init__(self, caminho):
        self.doc = Document(str(caminho))
        self.body = self.doc.element.body

    # ---------- busca ----------
    def achar(self, inicio: str | None = None, contem: str | None = None, estilo: str | None = None,
              apos=None) -> Paragraph:
        iniciar = apos is None
        for p in self.doc.paragraphs:
            if not iniciar:
                iniciar = p._p is apos
                continue
            if estilo and p.style.name != estilo:
                continue
            t = p.text.strip()
            if inicio and not t.startswith(inicio):
                continue
            if contem and contem not in t:
                continue
            return p
        raise LookupError(f"Trecho do template não encontrado: {inicio or contem!r}")

    # ---------- edição de texto ----------
    @staticmethod
    def cor(run) -> str:
        c = run._r.find(qn("w:rPr") + "/" + qn("w:color"))
        return (c.get(qn("w:val")) or "").lower() if c is not None else ""

    @classmethod
    def pintar_preto(cls, run):
        c = run._r.find(qn("w:rPr") + "/" + qn("w:color"))
        if c is not None:
            c.set(qn("w:val"), "000000")

    @classmethod
    def definir_run(cls, run, texto: str):
        run.text = texto
        if cls.cor(run) in VERMELHOS:
            cls.pintar_preto(run)

    @classmethod
    def substituir(cls, par: Paragraph, padrao: str, novo: str, flags=0) -> bool:
        """Regex dentro de cada run (preserva a formatação). Placeholders vermelhos viram texto preto."""
        mudou = False
        for r in par.runs:
            if r.text and re.search(padrao, r.text, flags):
                vermelho = cls.cor(r) in VERMELHOS
                r.text = re.sub(padrao, lambda m: novo, r.text, flags=flags)
                if vermelho:
                    cls.pintar_preto(r)
                mudou = True
        return mudou


    @classmethod
    def substituir_global(cls, par: Paragraph, padrao: str, novo: str, flags=0) -> bool:
        """Regex sobre o texto do parágrafo inteiro (pode cruzar runs). Preserva desenhos; o texto novo herda o run onde o trecho começa."""
        runs = [r for r in par.runs if r._r.find(qn("w:drawing")) is None and r._r.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent") is None]
        texto = "".join(r.text for r in runs)
        m = re.search(padrao, texto, flags)
        if not m:
            return False
        fatias, pos = [], 0
        for r in runs:
            fatias.append((r, pos, pos + len(r.text)))
            pos += len(r.text)
        ini, fim = m.span()
        primeiro = True
        for r, a, b in fatias:
            if b <= ini or a >= fim:
                continue
            antes = r.text[: max(ini - a, 0)]
            depois = r.text[max(fim - a, 0):] if fim < b else ""
            if primeiro:
                r.text = antes + novo + depois
                if cls.cor(r) in VERMELHOS:
                    cls.pintar_preto(r)
                primeiro = False
            else:
                r.text = depois
        return True

    @staticmethod
    def texto_unico(par: Paragraph, texto: str):
        """Coloca todo o texto no primeiro run com texto e esvazia os demais (mantém desenhos)."""
        usados = False
        for r in par.runs:
            if r._r.find(qn("w:drawing")) is not None or r._r.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent") is not None:
                continue
            if not usados and r.text.strip():
                r.text = texto
                usados = True
            elif r.text:
                r.text = ""
        if not usados and par.runs:
            par.runs[0].text = texto

    # ---------- estrutura ----------
    @staticmethod
    def remover(el):
        el.getparent().remove(el)

    def remover_entre(self, inicio, fim):
        """Remove os irmãos entre dois elementos (exclusivo)."""
        atual = inicio.getnext()
        while atual is not None and atual is not fim:
            prox = atual.getnext()
            self.remover(atual)
            atual = prox

    @staticmethod
    def tirar_secao(par_el):
        """Remove a quebra de seção (sectPr) de um parágrafo, mantendo o texto."""
        ppr = par_el.find(qn("w:pPr"))
        if ppr is not None:
            s = ppr.find(qn("w:sectPr"))
            if s is not None:
                ppr.remove(s)

    def novo_paragrafo(self, modelo=None):
        """Parágrafo novo (no fim do corpo) ou clone de um modelo; use Inserter para posicionar."""
        if modelo is not None:
            el = clone(modelo._p if isinstance(modelo, Paragraph) else modelo)
            for r in el.findall(qn("w:r")):
                el.remove(r)
            for h in el.findall(qn("w:hyperlink")):
                el.remove(h)
            ppr = el.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                ppr.remove(ppr.find(qn("w:sectPr")))
            self.body.insert(len(self.body) - 1, el)
            return Paragraph(el, self.doc._body)
        return self.doc.add_paragraph()

    def paragrafo_texto(self, modelo: Paragraph, runs: list[tuple[str, dict]] | str) -> Paragraph:
        """Clone de ``modelo`` com runs novos: [(texto, {'negrito': bool})] herdando fonte/tamanho do modelo."""
        p = self.novo_paragrafo(modelo)
        if isinstance(runs, str):
            runs = [(runs, {})]
        base = next((r for r in modelo.runs if r.text.strip() and r._r.find(qn("w:drawing")) is None), None)
        for txt, fmt in runs:
            r = p.add_run(txt)
            if base is not None and base._r.find(qn("w:rPr")) is not None:
                r._r.insert(0, clone(base._r.find(qn("w:rPr"))))
            if fmt.get("negrito") is not None:
                r.bold = fmt["negrito"]
            if fmt.get("italico") is not None:
                r.italic = fmt["italico"]
        return p

    def legenda(self, modelo: Paragraph, rotulo: str, texto: str, manter_com_proximo=True) -> Paragraph:
        """Legenda no padrão do template: rótulo em negrito + texto (clone de 'Tabela 02 - ...')."""
        p = self.novo_paragrafo(modelo)
        base = modelo.runs
        b = p.add_run(rotulo + " ")
        t = p.add_run("- " + texto)
        for novo, orig in ((b, base[0]), (t, base[1] if len(base) > 1 else base[0])):
            rpr = orig._r.find(qn("w:rPr"))
            if rpr is not None:
                novo._r.insert(0, clone(rpr))
        p.paragraph_format.keep_with_next = manter_com_proximo
        return p

    def paragrafo_imagem(self, fonte, largura_cm: float, alinhamento="center", altura_max_cm: float | None = None,
                         manter_com_proximo=False) -> Paragraph:
        p = self.doc.add_paragraph()
        p.alignment = {"center": 1, "left": 0}[alinhamento]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = manter_com_proximo
        run = p.add_run()
        pic = run.add_picture(fonte, width=Cm(largura_cm))
        if altura_max_cm and pic.height > Cm(altura_max_cm):
            esc = Cm(altura_max_cm) / pic.height
            pic.height, pic.width = int(pic.height * esc), int(pic.width * esc)
        return p

    @staticmethod
    def depois(ref, *elementos):
        """Insere elementos (ou objetos com ._p/._tbl) em sequência logo após ``ref``."""
        atual = ref
        for e in elementos:
            el = getattr(e, "_p", None) or getattr(e, "_tbl", None) or e
            atual.addnext(el)
            atual = el
        return atual

    # ---------- tabelas ----------
    @staticmethod
    def tabela(el) -> Table:
        return Table(el, None)


# ---------- formatação de células ----------
def sombrear(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def alinhar_vertical(cell, valor="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(old)
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), valor)
    tcPr.append(v)


def bordas_celula(cell, **lados):
    """bordas_celula(cell, bottom=('single', 6, '000000'), top=None ...)"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for lado, cfg in lados.items():
        e = b.find(qn(f"w:{lado}"))
        if e is None:
            e = OxmlElement(f"w:{lado}")
            b.append(e)
        val, sz, cor = cfg
        e.set(qn("w:val"), val); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), cor)


def margens_tabela(table, cm=0.08):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblCellMar")):
        tblPr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(int(cm * 567))); e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def bordas_tabela(table, cor="808080", sz=4):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), cor)
        b.append(e)
    tblPr.append(b)


def layout_fixo(table, larguras_cm: list[float], centralizar=True):
    table.autofit = False
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    w = OxmlElement("w:tblW"); w.set(qn("w:w"), str(int(sum(larguras_cm) * 567))); w.set(qn("w:type"), "dxa")
    tblPr.insert(1, w)
    if centralizar:
        for old in tblPr.findall(qn("w:jc")):
            tblPr.remove(old)
        for old in tblPr.findall(qn("w:tblInd")):
            tblPr.remove(old)
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
    for col, w_cm in zip(table.columns, larguras_cm):
        col.width = Cm(w_cm)
    for row in table.rows:
        cells = row.cells
        if len({id(c._tc) for c in cells}) != len(cells):     # linha com células mescladas
            continue
        for c, w_cm in zip(cells, larguras_cm):
            c.width = Cm(w_cm)


def linha_nao_quebra(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))


def linha_cabecalho(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def escrever(cell, texto, negrito=False, tam=7, cor=None, alinh="center", fonte="Arial", italico=False,
             antes=1, depois=1):
    """Substitui todo o conteúdo da célula por um único parágrafo formatado."""
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    for r in list(p._p):
        if r.tag != qn("w:pPr"):
            p._p.remove(r)
    ppr = p._p.get_or_add_pPr()
    for tag in ("w:spacing", "w:ind", "w:jc"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(antes * 20)); sp.set(qn("w:after"), str(depois * 20)); sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto")
    ppr.append(sp)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), {"center": "center", "left": "left", "right": "right"}[alinh]); ppr.append(jc)
    r = p.add_run(str(texto))
    r.bold, r.italic = negrito, italico
    r.font.size = Pt(tam)
    r.font.name = fonte
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:cs", "w:eastAsia"):
            rf.set(qn(a), fonte)
    if cor:
        from docx.shared import RGBColor
        r.font.color.rgb = RGBColor.from_string(cor)
    return p


def reescrever_celula_modelo(cell, texto, cor=None, negrito=None):
    """Troca o texto de uma célula copiada do template mantendo a fonte do parágrafo de texto.

    Remove os parágrafos vazios usados no template para centralizar na vertical e usa vAlign."""
    paras = cell.paragraphs
    com_texto = [p for p in paras if p.text.strip()] or [paras[-1]]
    alvo = com_texto[-1]
    for p in paras:
        if p is not alvo:
            p._p.getparent().remove(p._p)
    runs = [r for r in alvo.runs if r._r.find(qn("w:drawing")) is None]
    if runs:
        runs[0].text = str(texto)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        r0 = runs[0]
    else:
        r0 = alvo.add_run(str(texto))
    if cor:
        from docx.shared import RGBColor
        r0.font.color.rgb = RGBColor.from_string(cor)
    if negrito is not None:
        r0.bold = negrito
    ppr = alvo._p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is not None:
        sp.set(qn("w:before"), "20"); sp.set(qn("w:after"), "20")
    alinhar_vertical(cell, "center")
