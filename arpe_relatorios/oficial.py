"""Leitura dos blocos digitados à mão na aba 'Gráficos AAAA' (planilha oficial da CEEGC).

Nem tudo no Apêndice 2 pode ser deduzido das abas de dados: os *dias programados* e, principalmente,
o mês em que cada não conformidade (ou cliente) foi **sanado** são lançados manualmente nessa aba.
Este módulo lê esses valores para que o relatório reproduza exatamente o que a equipe acompanha."""
from __future__ import annotations

import unicodedata
from pathlib import Path

from .config import MESES_ABREV


def _norm(txt) -> str:
    s = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return " ".join(s.casefold().replace(":", " ").split())


def _numero(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


class PlanilhaOficial:
    """Extrai de 'Gráficos AAAA': meta anual, dias programados, NC sanadas e clientes sanados (por mês)."""

    def __init__(self, caminho: Path):
        self.caminho = Path(caminho)

    @staticmethod
    def nomes_aba(ano: int) -> list[str]:
        return [f"Gráficos {ano}", f"Graficos {ano}", f"Gráficos{ano}", f"Gráfico {ano}"]

    def ler(self, ano: int) -> dict:
        """Devolve {} se a aba não existir. Séries são listas de 12 posições (None = célula vazia)."""
        import openpyxl
        wb = openpyxl.load_workbook(self.caminho, read_only=True, data_only=True)
        try:
            aba = next((n for n in self.nomes_aba(ano) if n in wb.sheetnames), None)
            if not aba:
                return {}
            linhas = [list(r) for r in wb[aba].iter_rows(values_only=True)]
        finally:
            wb.close()
        return self._extrair(linhas)

    # ------------------------------------------------------------------
    @staticmethod
    def _colunas_meses(linhas: list[list], i_label: int) -> dict[int, int]:
        """Procura, acima do rótulo, a linha de cabeçalho com JAN..DEZ e devolve {mes(0-11): coluna}."""
        for i in range(i_label - 1, max(i_label - 8, -1), -1):
            cols = {}
            for j, v in enumerate(linhas[i]):
                n = _norm(v).upper()
                if n in MESES_ABREV and MESES_ABREV.index(n) not in cols:
                    cols[MESES_ABREV.index(n)] = j
            if len(cols) >= 6:
                return cols
        return {}

    def _serie(self, linhas, i, cols) -> list:
        return [(_numero(linhas[i][cols[m]]) if m in cols and cols[m] < len(linhas[i]) else None) for m in range(12)]

    def _extrair(self, linhas: list[list]) -> dict:
        out: dict = {}
        for i, row in enumerate(linhas):
            rotulos = [_norm(c) for c in row[:5] if isinstance(c, str)]
            if not rotulos:
                continue
            for rot in rotulos:
                if rot.startswith("meta anual") and "meta_anual" not in out:
                    nums = [_numero(c) for c in row if _numero(c) is not None]
                    if nums:
                        out["meta_anual"] = int(nums[0])
                chave = None
                if rot.startswith("dias programados"):
                    chave = "dias_programados"
                elif rot.startswith("nc sanadas"):
                    chave = "nc_sanadas"
                elif rot.split(" ")[0] in ("estacoes", "clientes", "usuarios") and "sanad" in rot:
                    chave = "clientes_sanados"
                if chave and chave not in out:
                    cols = self._colunas_meses(linhas, i)
                    if cols:
                        out[chave] = self._serie(linhas, i, cols)
        return out
