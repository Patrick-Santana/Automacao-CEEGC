from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from .analytics import DashboardAnalytics
from .charts import ChartFactory
from .config import Config
from .models import Dataset, Fiscalizacao, Periodo
from .photos import FotoItem, PhotoLibrary, visita_ids

AZUL_CAB = "D9E2F3"


def _lista_pt(itens: list[str]) -> str:
    itens = [i for i in itens if i]
    return itens[0] if len(itens) == 1 else (", ".join(itens[:-1]) + " e " + itens[-1] if itens else "")


class DocHelper:
    """Pequenos utilitários de formatação python-docx."""

    @staticmethod
    def sombra(cell, hex_):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
        tcPr.append(shd)

    @staticmethod
    def margens_celula(table, cm=0.12):
        tblPr = table._tbl.tblPr
        mar = OxmlElement("w:tblCellMar")
        for lado in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:w"), str(int(cm * 567))); e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)

    @staticmethod
    def linha_nao_quebra(row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

    @staticmethod
    def cabecalho_repete(row):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true")
        trPr.append(e)

    @staticmethod
    def larguras(table, cms):
        table.autofit = False
        tblPr = table._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
        tblPr.append(lay)
        for col, w in zip(table.columns, cms):
            col.width = Cm(w)
        for row in table.rows:
            cells = row.cells
            if len({id(c._tc) for c in cells}) != len(cells):   # linha com células mescladas
                continue
            for c, w in zip(cells, cms):
                c.width = Cm(w)

    @staticmethod
    def texto(cell, txt, negrito=False, tam=9, cor=None, alinh=WD_ALIGN_PARAGRAPH.CENTER, italico=False):
        p = cell.paragraphs[0]
        p.alignment = alinh
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run(str(txt))
        r.bold, r.italic, r.font.size = negrito, italico, Pt(tam)
        if cor:
            r.font.color.rgb = RGBColor.from_string(cor)
        return p


class PhotoTableBuilder:
    """Insere fotos em tabelas com borda: até 6 por tabela (3 linhas × 2 colunas) + legenda."""

    LARG_MAX, ALT_MAX = 7.6, 6.1   # cm por célula (imagem)

    def __init__(self, config: Config):
        self.cfg = config

    def _tamanho(self, path: Path) -> tuple[float, float]:
        with Image.open(path) as im:
            w, h = im.size
        esc = min(self.LARG_MAX / w, self.ALT_MAX / h)
        return w * esc, h * esc

    def adicionar(self, doc: Document, itens: list[FotoItem], numerar_desde: int = 1):
        n_tab = self.cfg.fotos_por_tabela
        cols = self.cfg.fotos_colunas
        for bloco_i in range(0, len(itens), n_tab):
            bloco = itens[bloco_i:bloco_i + n_tab]
            if bloco_i > 0:
                doc.add_page_break()
            linhas = -(-len(bloco) // cols)
            t = doc.add_table(rows=linhas, cols=cols)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            DocHelper.margens_celula(t, 0.15)
            for k, item in enumerate(bloco):
                cell = t.cell(k // cols, k % cols)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                w, h = self._tamanho(item.path)
                p.add_run().add_picture(str(item.path), width=Cm(w), height=Cm(h))
                cap = cell.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(6)
                r = cap.add_run(f"Foto {numerar_desde + bloco_i + k}: {item.legenda}")
                r.italic, r.font.size = True, Pt(9)
            for row in t.rows:
                DocHelper.linha_nao_quebra(row)
            DocHelper.larguras(t, [8.2] * cols)


class ReportBuilder:
    """Monta o relatório mensal (.docx) seguindo a estrutura do modelo da ARPE."""

    def __init__(self, config: Config, dataset: Dataset):
        self.cfg, self.ds = config, dataset
        self.photos = PhotoLibrary(config)
        self.fotos_tab = PhotoTableBuilder(config)

    # ---------- helpers de texto ----------
    def _h(self, doc, txt, nivel=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if nivel == 1 else 8)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(txt)
        r.bold, r.font.size = True, Pt(12 if nivel == 1 else 11)
        return p

    def _p(self, doc, txt, negrito_ini: str | None = None, recuo=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if recuo:
            p.paragraph_format.first_line_indent = Cm(1.25)
        if negrito_ini:
            p.add_run(negrito_ini).bold = True
        p.add_run(txt)
        return p

    def _legenda_tabela(self, doc, txt):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        b, _, resto = txt.partition(" - ")
        p.add_run(b + " ").bold = True
        p.add_run("- " + resto)
        for r in p.runs:
            r.font.size = Pt(9)

    def _df_tabela(self, doc, df: pd.DataFrame, indice_titulo="", largura_total=17.0, tam=8):
        cols = list(df.columns)
        t = doc.add_table(rows=1, cols=len(cols) + 1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        DocHelper.margens_celula(t, 0.06)
        for i, h in enumerate([indice_titulo] + [str(c) for c in cols]):
            DocHelper.texto(t.rows[0].cells[i], h, True, tam)
            DocHelper.sombra(t.rows[0].cells[i], AZUL_CAB)
        DocHelper.cabecalho_repete(t.rows[0])
        for idx, row in df.iterrows():
            cells = t.add_row().cells
            DocHelper.texto(cells[0], idx, str(idx).lower() == "total", tam, alinh=WD_ALIGN_PARAGRAPH.LEFT)
            for j, v in enumerate(row, start=1):
                if isinstance(v, float) and pd.isna(v):
                    s = ""
                elif isinstance(v, float) and v == int(v) and cols[j - 1] != "%":
                    s = str(int(v))
                elif isinstance(v, float):
                    s = f"{v:.2f}".replace(".", ",")
                else:
                    s = str(v)
                DocHelper.texto(cells[j], s, str(idx).lower() == "total", tam)
        primeira = 3.6 if len(cols) > 8 else 5.0
        resto = (largura_total - primeira) / len(cols)
        DocHelper.larguras(t, [primeira] + [resto] * len(cols))
        return t

    def _fig(self, doc, fig, largura=15.5):
        doc.add_picture(ChartFactory.para_png(fig), width=Cm(largura))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- construção ----------
    def build(self, p: Periodo, destino: Path | None = None) -> Path:
        fiscs = self.ds.do_periodo(p)
        if not fiscs:
            raise ValueError(f"Não há fiscalizações em {p.label}.")
        vids = visita_ids(fiscs)
        an = DashboardAnalytics(self.ds, p.ano, self.cfg)
        ncs = [(f, nc) for f in fiscs for nc in f.ncs]
        cidades = list(vids)
        analistas = sorted({n.strip() for f in fiscs for n in
                            f.func_arpe.replace(" e ", ",").replace("/", ",").split(",") if n.strip()})

        doc = Document()
        sec = doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        st = doc.styles["Normal"]
        st.font.name, st.font.size = "Calibri", Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

        self._capa(doc, p, cidades, analistas)
        self._intro(doc, p, an, cidades)
        self._fiscalizacao(doc, p, an, fiscs, ncs, vids)
        self._determinacoes(doc, ncs)
        self._apendice1(doc, fiscs, ncs, vids)
        self._apendice2(doc, an, p)
        self._assinaturas(doc, analistas)

        destino = destino or self.cfg.saida_dir / f"Relatorio_Fiscalizacao_{p.ano}-{p.mes:02d}.docx"
        destino.parent.mkdir(parents=True, exist_ok=True)
        doc.save(destino)
        return destino

    def _capa(self, doc, p, cidades, analistas):
        def centro(txt, tam=12, neg=True, antes=0, cor=None):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(antes)
            r = par.add_run(txt)
            r.bold, r.font.size = neg, Pt(tam)
            if cor:
                r.font.color.rgb = RGBColor.from_string(cor)
        centro("AGÊNCIA DE REGULAÇÃO DE PERNAMBUCO – ARPE", 14, True, 30, "1F3864")
        centro("Coordenadoria de Energia Elétrica e Gás Canalizado – CEEGC", 11, False)
        centro("RELATÓRIO DE FISCALIZAÇÃO", 20, True, 170)
        centro(f"FISCALIZAÇÃO DAS INSTALAÇÕES DE GÁS NO(S) MUNICÍPIO(S) DE {_lista_pt(cidades).upper()}", 13, True, 40)
        for a in analistas:
            centro(a, 12, True, 4 if a != analistas[0] else 60)
        centro(f"{p.nome_mes.upper()} {p.ano}", 12, True, 6)
        centro("RELATÓRIO DE FISCALIZAÇÃO DIRETA – PROCESSO ADMINISTRATIVO", 11, True, 90)
        centro(f"PA-{p.mes:03d}/{p.ano}-CEEGC – GÁS PROCESSOS", 11, True)
        centro("SEI Nº _______________________", 11, True, 0, "C00000")
        doc.add_page_break()
        self._h(doc, "LISTA DE ABREVIATURAS E SIGLAS")
        siglas = [("CRM", "Conjunto de Regulagem de Pressão e Medição"), ("ERP", "Estação de Regulagem e Pressão"),
                  ("ERPM", "Estação de Regulagem, Pressão e Medição"), ("ETC", "Estação de Transferência de Custódia"),
                  ("GN", "Gás Natural"), ("GNL", "Gás Natural Liquefeito"), ("GNV", "Gás Natural Veicular"),
                  ("PSV", "Válvula de Segurança e Alívio de Pressão"), ("MVAZ", "Medidor de Vazão")]
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        DocHelper.margens_celula(t, 0.1)
        for c, h in zip(t.rows[0].cells, ("Sigla", "Definição")):
            DocHelper.texto(c, h, True, 10); DocHelper.sombra(c, AZUL_CAB)
        for s, d in siglas:
            c = t.add_row().cells
            DocHelper.texto(c[0], s, True, 10); DocHelper.texto(c[1], d, False, 10, alinh=WD_ALIGN_PARAGRAPH.LEFT)
        DocHelper.larguras(t, [3, 14])
        doc.add_page_break()

    def _intro(self, doc, p, an, cidades):
        lst = _lista_pt(cidades)
        self._h(doc, "1. INTRODUÇÃO")
        self._p(doc, "Atualmente, a prestação dos serviços públicos de odorização, canalização e distribuição de gás natural em "
                     "Pernambuco é realizada pela Companhia Pernambucana de Gás (Copergás). A Agência de Regulação dos Serviços "
                     "Públicos Delegados do Estado de Pernambuco (Arpe), por meio da Coordenadoria de Energia Elétrica e Gás "
                     "Canalizado (CEEGC), conduz fiscalizações voltadas à regulação técnico-operacional dos serviços prestados, "
                     "avaliando as condições operacionais, a conservação e a manutenção das instalações, a conformidade com a "
                     "legislação vigente, a qualidade do serviço e a satisfação dos usuários.")
        self._p(doc, "As Fiscalizações Periódicas, organizadas na Agenda Regulatória da CEEGC, inspecionam se as instalações do "
                     f"sistema de distribuição atendem às normas legais. Este relatório apresenta os resultados das fiscalizações "
                     f"realizadas in loco em {lst}, durante o mês de {p.nome_mes} de {p.ano}.")
        self._h(doc, "2. OBJETIVOS")
        self._p(doc, "A fiscalização direta e periódica verifica o grau de conformidade das unidades operacionais com a legislação "
                     "e as normas vigentes e determina ou recomenda medidas corretivas. Objetivos específicos:")
        self._p(doc, " verificar o cumprimento das normas aplicáveis ao gás canalizado, em especial nos equipamentos dos sistemas "
                     "de distribuição (medidores, lacres, tubulações e placas de identificação);", "Conformidade Legal:", False)
        self._p(doc, " analisar as condições técnico-operacionais, o estado de conservação, a manutenção e a segurança das unidades.",
                "Condições Operacionais, de Conservação e Manutenção:", False)
        self._p(doc, f"De acordo com a Agenda Regulatória, a meta de fiscalizações para {p.ano} é de {an.meta_anual}. Considerando essa "
                     "meta, a equipe procura abranger todos os municípios com instalações da Copergás e os diferentes nichos.")
        self._h(doc, "3. METODOLOGIA")
        self._p(doc, f"A fiscalização direta e periódica em {lst} é realizada por analistas da CEEGC e organizada em três etapas: "
                     "Preparação e Planejamento, Execução da Fiscalização e Monitoramento e Avaliação. A execução é pautada pelo "
                     "seguinte arcabouço normativo:")
        for n in ("Lei Federal nº 14.134/2021 (atividades relativas ao transporte de gás natural);",
                  "Lei Estadual nº 17.641/2022, que altera a Lei Estadual nº 15.900/2016 (serviços locais de gás canalizado em Pernambuco);",
                  "Lei Estadual nº 12.524/2003, que consolida as disposições sobre a criação da Arpe (arts. 3º e 4º);",
                  "Resolução Arpe nº 034/2006 (prestação do serviço de fornecimento de gás canalizado, indicadores e penalidades);",
                  "Resolução Arpe nº 083/2013 (procedimentos de fiscalização, autuação e aplicação de penalidades);",
                  "ABNT NBR 12.712 (projeto de sistemas de transmissão e distribuição de gás combustível);",
                  "ABNT NBR 15.526 (redes de distribuição interna para gases combustíveis em instalações residenciais)."):
            par = doc.add_paragraph(n, style="List Bullet")
            par.paragraph_format.space_after = Pt(2)

    def _fiscalizacao(self, doc, p, an, fiscs, ncs, vids):
        self._h(doc, "4. FISCALIZAÇÃO")
        self._h(doc, "4.1. Preparação e Planejamento", 2)
        self._p(doc, f"A equipe da CEEGC definiu o(s) município(s) de {_lista_pt(list(vids))} para as fiscalizações do mês de "
                     f"{p.nome_mes}. A Copergás foi informada e designou funcionário para acompanhar cada dia de fiscalização.")
        acum = int((an.f.mes <= p.mes).sum())
        pct = str(round(100 * acum / an.meta_anual, 1)).replace(".", ",")
        self._p(doc, f"Até {p.nome_mes} de {p.ano} foram realizadas {acum} fiscalizações, o que corresponde a {pct}% "
                     f"da meta anual de {an.meta_anual}.")

        self._h(doc, "4.2. Execução da Fiscalização", 2)
        self._p(doc, "As tabelas a seguir apresentam os clientes fiscalizados em cada dia, com os membros da equipe da CEEGC e o "
                     "funcionário da Copergás que acompanhou. As Não Conformidades (NC) constatadas in loco estão relacionadas na "
                     "Tabela de Não Conformidades, e seus registros fotográficos estão no Apêndice 1.")
        n_tab = 1
        dias = sorted({f.data for f in fiscs})
        for dia in dias:
            do_dia = [f for f in fiscs if f.data == dia]
            arpe = _lista_pt(sorted({f.func_arpe for f in do_dia if f.func_arpe}))
            cop = _lista_pt(sorted({f.func_copergas for f in do_dia if f.func_copergas}))
            self._legenda_tabela(doc, f"Tabela {n_tab:02d} - Fiscalizações realizadas no dia {dia.strftime('%d/%m/%Y')}")
            n_tab += 1
            t = doc.add_table(rows=2, cols=4)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            DocHelper.margens_celula(t, 0.08)
            for i, txt in enumerate((f"Equipe Arpe: {arpe}", f"Funcionário Copergás: {cop}")):
                c = t.rows[i].cells[0].merge(t.rows[i].cells[3])
                DocHelper.texto(c, txt, True, 9, alinh=WD_ALIGN_PARAGRAPH.LEFT)
            hdr = t.add_row().cells
            for c, h in zip(hdr, ("Nº da Fiscalização", "Cliente", "Endereço", "Status")):
                DocHelper.texto(c, h, True, 9); DocHelper.sombra(c, AZUL_CAB)
            DocHelper.cabecalho_repete(t.rows[2])
            for f in do_dia:
                c = t.add_row().cells
                DocHelper.texto(c[0], f.seq, False, 9)
                DocHelper.texto(c[1], f.usuario, False, 9)
                DocHelper.texto(c[2], f.endereco_completo, False, 9)
                DocHelper.texto(c[3], "CONFORME" if f.conforme else "NÃO CONFORME", True, 8,
                                "2E7D32" if f.conforme else "C00000")
                DocHelper.linha_nao_quebra(t.rows[-1])
            DocHelper.larguras(t, [2.2, 5.6, 6.7, 2.5])

        self._legenda_tabela(doc, f"Tabela {n_tab:02d} - Relação das Não Conformidades constatadas")
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        DocHelper.margens_celula(t, 0.08)
        for c, h in zip(t.rows[0].cells, ("Fiscalização Nº", "Usuário", "Não conformidade", "Registro fotográfico",
                                          "Fundamentação legal (Res. Arpe 34/2006)", "Qtd. NC")):
            DocHelper.texto(c, h, True, 8); DocHelper.sombra(c, AZUL_CAB)
        DocHelper.cabecalho_repete(t.rows[0])
        foto_n = 1
        for f, nc in ncs:
            n_fotos = max(1, len(self.photos.fotos_nc(nc, vids[f.cidade], f.cidade)))
            refs = ", ".join(f"Foto {foto_n + k:02d}" for k in range(n_fotos))
            foto_n += n_fotos
            c = t.add_row().cells
            for cell, txt, al in ((c[0], f.seq, None), (c[1], f.usuario, None), (c[2], nc.detalhe, None),
                                  (c[3], refs, None), (c[4], self.cfg.fundamentacao_padrao, None), (c[5], 1, None)):
                DocHelper.texto(cell, txt, False, 8)
            DocHelper.linha_nao_quebra(t.rows[-1])
        c = t.add_row().cells
        DocHelper.texto(c[0], "TOTAL", True, 8)
        DocHelper.texto(c[5], len(ncs), True, 8)
        DocHelper.larguras(t, [1.9, 4.0, 4.6, 2.2, 3.2, 1.3])
        if not ncs:
            self._p(doc, "Não foram constatadas não conformidades nas fiscalizações do período.", recuo=False)
        else:
            self._p(doc, "")
            self._p(doc, "As Não Conformidades constatadas caracterizam descumprimento das normas vigentes e/ou do contrato de "
                         "concessão. Conforme o artigo 35 da Resolução Arpe nº 034/2006:")
            self._p(doc, "“Art. 35 - As infrações às disposições legais e contratuais relativas à prestação de serviços, implantação "
                         "e operação de instalações de distribuição de gás canalizado ou serviços autorizados sujeitarão a "
                         "concessionária às penalidades de: I - Advertência; II - Multa; e, III - Intervenção administrativa.”",
                    recuo=False).paragraph_format.left_indent = Cm(3)
            self._p(doc, "As infrações foram enquadradas no inciso V do artigo 38 da mesma Resolução: “Art. 38 - Constitui infração, "
                         "sujeita à imposição da penalidade de multa do Tipo II, o fato de: (...) V - não acatar as normas técnicas e "
                         "recomendações estabelecidas para projetos, construção, operação e manutenção das instalações de "
                         "distribuição de gás canalizado, nos termos do Contrato de Concessão e da legislação;”",
                    recuo=False).paragraph_format.left_indent = Cm(3)
        self._h(doc, "4.3. Monitoramento e Avaliação", 2)
        self._p(doc, "Após a execução da fiscalização in loco, seguem os trâmites das Resoluções Arpe nº 34/2006 e nº 83/2013: "
                     "Determinações, Recomendações, Termo de Notificação, Relatório de Fiscalização, Plano de Ação, Relatórios de "
                     "Acompanhamento e Avaliação Final.")

    def _determinacoes(self, doc, ncs):
        self._h(doc, "5. DETERMINAÇÕES GERAIS")
        if ncs:
            self._p(doc, "Diante das constatações apontadas no presente relatório, recomenda-se que a Copergás adote providências "
                         "quanto às não conformidades mencionadas, a fim de atender à Resolução Arpe nº 034, de 10 de agosto de 2006, "
                         "bem como ao contrato de concessão e demais normas pertinentes. Reiteramos a importância do monitoramento "
                         "contínuo e da manutenção dos sistemas de distribuição de gás natural para a eficiência dos serviços prestados.")
        else:
            self._p(doc, "Não foram constatadas não conformidades no período. Reiteramos a importância do monitoramento contínuo e da "
                         "manutenção dos sistemas de distribuição de gás natural para a eficiência dos serviços prestados.")

    def _apendice1(self, doc, fiscs, ncs, vids):
        doc.add_page_break()
        self._h(doc, "APÊNDICE 1 - FOTOS DAS NÃO CONFORMIDADES")
        itens: list[FotoItem] = []
        for f, nc in ncs:
            for path in self.photos.fotos_nc(nc, vids[f.cidade], f.cidade):
                itens.append(FotoItem(path, f"Fiscalização Nº {f.seq} - {nc.detalhe}"))
        if itens:
            self.fotos_tab.adicionar(doc, itens)
        else:
            self._p(doc, "Nenhum registro fotográfico de não conformidade disponível para o período.", recuo=False)
        gerais = []
        for cidade, vid in vids.items():
            for i, path in enumerate(self.photos.fotos_gerais(vid, cidade), start=1):
                gerais.append(FotoItem(path, f"Condições gerais – {cidade}"))
        if gerais:
            doc.add_page_break()
            self._h(doc, "APÊNDICE 1-A - REGISTROS DAS CONDIÇÕES GERAIS")
            self.fotos_tab.adicionar(doc, gerais)

    def _apendice2(self, doc, an, p):
        doc.add_page_break()
        self._h(doc, "APÊNDICE 2 - ANÁLISE DAS FISCALIZAÇÕES")
        k = an.kpis()
        pct = str(round(k["pct_meta"], 1)).replace(".", ",")
        self._p(doc, "Com o intuito de aprimorar a ação regulatória e a efetividade dos serviços públicos prestados pela "
                     "concessionária, apresentamos a evolução das fiscalizações do período e o acompanhamento das não conformidades.")
        self._p(doc, f"Em {p.ano}, até o momento, foram realizadas {k['realizado']} fiscalizações em {k['municipios']} municípios "
                     f"({pct}% da meta anual de {k['meta']}). Foram detectadas {k['nc_detectadas']} não conformidades, das quais "
                     f"{k['nc_pendentes']} permanecem pendentes.")
        seg = an.por_segmento_mes()
        self._legenda_tabela(doc, "Tabela A - Fiscalizações por segmento e mês")
        self._df_tabela(doc, seg, "Segmento")
        self._legenda_tabela(doc, "Tabela B - Fiscalizações por município e mês")
        self._df_tabela(doc, an.por_cidade_mes(), "Município", tam=7)
        self._legenda_tabela(doc, "Gráfico 01 - Quantitativo de fiscalizações por mês")
        self._fig(doc, ChartFactory(an).mensal())
        self._legenda_tabela(doc, "Gráfico 02 - Evolução acumulada das fiscalizações")
        self._fig(doc, ChartFactory(an).acumulado())
        self._legenda_tabela(doc, "Tabela C - Dias utilizados e média de fiscalizações por dia")
        self._df_tabela(doc, an.dias_e_media(), "")
        self._legenda_tabela(doc, "Tabela D - Acompanhamento das não conformidades")
        self._df_tabela(doc, an.acompanhamento_nc(), "")
        self._legenda_tabela(doc, "Gráfico 03 - Acompanhamento mensal das não conformidades")
        self._fig(doc, ChartFactory(an).nc_evolucao())
        self._legenda_tabela(doc, "Tabela E - Não conformidades mais frequentes")
        fr = an.frequencia_nc()
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, h in zip(t.rows[0].cells, ("Não conformidade", "Frequência", "%")):
            DocHelper.texto(c, h, True, 9); DocHelper.sombra(c, AZUL_CAB)
        for _, r in fr.iterrows():
            c = t.add_row().cells
            DocHelper.texto(c[0], r["Não conformidade"], False, 9, alinh=WD_ALIGN_PARAGRAPH.LEFT)
            DocHelper.texto(c[1], int(r["Frequência"]), False, 9)
            DocHelper.texto(c[2], f"{r['%']:.2f}".replace(".", ","), False, 9)
        c = t.add_row().cells
        DocHelper.texto(c[0], "Total", True, 9); DocHelper.texto(c[1], int(fr["Frequência"].sum()), True, 9)
        DocHelper.texto(c[2], "100,00", True, 9)
        DocHelper.larguras(t, [11.0, 3.0, 3.0])

    def _assinaturas(self, doc, analistas):
        doc.add_paragraph()
        par = doc.add_paragraph("Recife, data da assinatura eletrônica.")
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for a in analistas:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(14)
            par.add_run(a).bold = True
            par.add_run("\nAnalista de Regulação e Fiscalização\nMatrícula ____________")
