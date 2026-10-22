"""Gerador do relatório mensal: preenche o Template.docx da ARPE (em vez de montar o documento do zero).

O que é mantido do template: capa, siglas, sumário, introdução, objetivos, metodologia, textos legais,
determinações, assinaturas e rodapé.  O que é preenchido/gerado: campos vermelhos (XXXX), tabelas de meta por
município, tabelas de fiscalizações por dia, relação de não conformidades, fotos (Apêndice 1) e todo o Apêndice 2.
"""
from __future__ import annotations

import copy
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

from .analytics import ORDEM_SEGMENTOS, DashboardAnalytics
from .charts import ChartFactory
from .config import MESES_ABREV, MESES_NOME, Config
from .docx_tools import (Editor, bordas_celula, bordas_tabela, clone, escrever, layout_fixo, linha_cabecalho,
                         linha_nao_quebra, margens_tabela, reescrever_celula_modelo, sombrear)
from .metas import MetaMunicipios, normalizar
from .models import Dataset, Fiscalizacao, NaoConformidade, Periodo
from .narrativa import Narrativa, Numeracao, extenso, lista_pt, num, pct
from .photos import FotoItem, PhotoLibrary, SyntheticPhotoGenerator, visita_ids

AZUL_BANDA, CINZA_ZEBRA = "9DC3E6", "E7E6E6"
VERDE_OK, VERMELHO_NC = "385824", "FF0000"
MESES_UP = [m.upper() for m in MESES_NOME]


# ======================================================================================================
@dataclass
class Opcoes:
    """Dados que variam a cada relatório e não estão na planilha."""
    pa: str = ""                      # nº do processo administrativo (ex.: "05"); vazio = mantém o campo vermelho
    sei: str = ""                     # nº SEI; vazio = mantém o campo vermelho
    incluir_condicoes_gerais: bool = False
    gerar_fotos_sinteticas: bool = False


@dataclass
class Resultado:
    caminho: Path
    periodo: Periodo
    n_fiscalizacoes: int = 0
    n_nc: int = 0
    n_fotos: int = 0
    avisos: list[str] = field(default_factory=list)


def titulo_pt(nome: str) -> str:
    peq = {"de", "da", "do", "das", "dos", "e"}
    partes = nome.strip().lower().split()
    return " ".join(w if (i and w in peq) else w[:1].upper() + w[1:] for i, w in enumerate(partes))


def _data_br(d) -> str:
    return d.strftime("%d/%m/%Y")


# ======================================================================================================
class RelatorioBuilder:
    def __init__(self, cfg: Config, ds: Dataset):
        self.cfg, self.ds = cfg, ds
        self.metas = MetaMunicipios(cfg, ds)
        self.fotos = PhotoLibrary(cfg)
        self.avisos: list[str] = []

    # --------------------------------------------------------------------------------------------------
    def build(self, p: Periodo, destino: Path | None = None, opcoes: Opcoes | None = None) -> Resultado:
        op = opcoes or Opcoes()
        self.avisos = []
        fiscs = self.ds.do_periodo(p)
        if not fiscs:
            raise ValueError(f"Não há fiscalizações em {p.label} na planilha.")
        if op.gerar_fotos_sinteticas:
            SyntheticPhotoGenerator(self.cfg).gerar_periodo(self.ds, p, condicoes_gerais=op.incluir_condicoes_gerais)

        self.p, self.op, self.fiscs = p, op, fiscs
        self.an = DashboardAnalytics(self.ds, p.ano, self.cfg, ate_mes=p.mes)
        self.avisos += self.an.avisos()
        self.vids = visita_ids(fiscs)
        self.cidades = list(self.vids)                                   # na ordem da 1ª fiscalização
        self.ncs = [(f, nc) for f in fiscs for nc in sorted(f.ncs, key=lambda n: n.num)]
        self.nm = Numeracao(len({f.data for f in fiscs}))
        self.cli_nc_mes = len({f.id for f, _ in self.ncs})
        self.narr = Narrativa(self.cfg, self.ds, p, self.an, self.nm, self.cidades, len(self.ncs), self.cli_nc_mes)

        self.ed = Editor(self.cfg.template_efetivo())
        self._capturar_ancoras()
        self._preparar_fotos()
        self._capa()
        self._textos_fixos()
        self._secao_41()
        self._secao_42()
        self._secao_43_e_5()
        self._apendice1()
        self._apendice2()

        destino = Path(destino) if destino else self.cfg.saida_dir / f"Relatorio_Fiscalizacao_{p.ano}-{p.mes:02d}.docx"
        destino.parent.mkdir(parents=True, exist_ok=True)
        self.ed.doc.save(destino)
        return Resultado(destino, p, len(fiscs), len(self.ncs), len(self.itens_fotos), self.avisos)

    # --------------------------------------------------------------------------------------------------
    # Âncoras do template (localizadas por texto: sobrevivem a pequenas edições do modelo)
    # --------------------------------------------------------------------------------------------------
    def _capturar_ancoras(self):
        ed = self.ed
        a = {}
        a["capa_mun"] = ed.achar(inicio="FISCALIZAÇÃO DAS INSTALAÇÕES DE GÁS")
        a["an1"] = ed.achar(inicio="Analista 1")
        a["an2"] = ed.achar(inicio="Analista 2")
        a["pa"] = ed.achar(inicio="RELATÓRIO DE FISCALIZAÇÃO DIRETA")
        a["sei"] = ed.achar(inicio="SEI N")
        a["intro2"] = ed.achar(contem="mais recentes fiscalizações")
        a["meta"] = ed.achar(inicio="De acordo com a Agenda Regulatória da Coordenadoria")
        a["metod"] = ed.achar(inicio="A fiscalização direta e periódica do(s)")
        a["analise"] = ed.achar(inicio="Análise de Fiscalizações anteriores")
        a["para"] = ed.achar(inicio="Para isso, a meta foi distribuída")
        a["cap01"] = ed.achar(inicio="Tabela 01")
        a["fig01"] = ed.achar(inicio="Figura 01")
        a["def"] = ed.achar(inicio="Definição dos municípios")
        a["exec"] = ed.achar(inicio="As Tabelas 02 e 03 apresentam")
        a["cap_dia1"] = ed.achar(inicio="Tabela 02")
        a["cap_nc"] = ed.achar(inicio="Tabela 04", contem="Relação das Não Conformidades")
        a["nc_leg"] = ed.achar(inicio="As Não Conformidades constatadas", apos=a["cap_nc"]._p)
        a["h_monit"] = ed.achar(inicio="Monitoramento e Avaliação", apos=a["cap_nc"]._p)
        a["p43"] = ed.achar(inicio="Após a execução da fiscalização in loco", apos=a["h_monit"]._p)
        a["cap_meta2"] = ed.achar(contem="Meta por municípios em", apos=a["p43"]._p)
        a["deter"] = ed.achar(inicio="Diante das constatações")
        a["ap1"] = ed.achar(inicio="APÊNDICE 1", apos=a["deter"]._p)
        a["ap2"] = ed.achar(inicio="APÊNDICE 2", apos=a["ap1"]._p)
        a["ap2_intro"] = ed.achar(inicio="Com o intuito de aprimorar", apos=a["ap2"]._p)
        a["ap2_primeiro"] = ed.achar(inicio="As Tabelas 05 e 06", apos=a["ap2_intro"]._p)
        a["assin"] = ed.achar(inicio="Recife, data da assinatura", apos=a["ap2_primeiro"]._p)
        self.a = a

        def proximo_tbl(el):
            while el is not None and el.tag != qn("w:tbl"):
                el = el.getnext()
            return el
        self.t01a = proximo_tbl(a["cap01"]._p.getnext())
        self.t01b = proximo_tbl(self.t01a.getnext())
        self.t_dia = proximo_tbl(a["cap_dia1"]._p.getnext())
        self.t_nc = proximo_tbl(a["cap_nc"]._p.getnext())
        self.t02a = proximo_tbl(a["cap_meta2"]._p.getnext())
        self.t02b = proximo_tbl(self.t02a.getnext())
        # modelos (cópias) usados depois que o conteúdo original for removido
        self.m_txt = Paragraph(clone(a["ap2_intro"]._p), ed.doc._body)
        self.m_cap = Paragraph(clone(a["cap_dia1"]._p), ed.doc._body)
        self.m_cap_esq = Paragraph(clone(a["cap01"]._p), ed.doc._body)
        self.m_h = Paragraph(clone(a["ap1"]._p), ed.doc._body)
        self.tp_meta_a, self.tp_meta_b = clone(self.t01a), clone(self.t01b)
        self.tp_dia, self.tp_nc = clone(self.t_dia), clone(self.t_nc)
        self.tp_meta2a, self.tp_meta2b = clone(self.t02a), clone(self.t02b)

    # --------------------------------------------------------------------------------------------------
    # Fotos
    # --------------------------------------------------------------------------------------------------
    def _preparar_fotos(self):
        self.itens_fotos: list[FotoItem] = []
        self.refs_foto: dict[int, list[int]] = {}
        for f, nc in self.ncs:
            caminhos = self.fotos.fotos_nc(nc, f, self.vids[f.cidade])
            ini = len(self.itens_fotos) + 1
            for c in caminhos:
                self.itens_fotos.append(FotoItem(c, f"Fiscalização Nº {f.seq} - {nc.detalhe_limpo}"))
            self.refs_foto[nc.num] = list(range(ini, ini + len(caminhos)))
            if not caminhos:
                self.avisos.append(f"Sem foto para a NC {nc.num} (fiscalização {f.seq} – {nc.detalhe_limpo}).")

    # --------------------------------------------------------------------------------------------------
    # Capa
    # --------------------------------------------------------------------------------------------------
    def _nomes_equipe(self) -> list[str]:
        vistos: list[str] = []
        for f in self.fiscs:
            for n in f.equipe_arpe:
                if n.casefold() not in [v.casefold() for v in vistos]:
                    vistos.append(n)
        completos = {a["nome"].split()[0].casefold(): a["nome"] for a in self.cfg.assinantes}
        faltam = [n for n in vistos if n.split()[0].casefold() not in completos]
        if faltam:
            self.avisos.append("Nome completo não cadastrado (aparece só o primeiro nome na capa): "
                               + ", ".join(faltam) + ". Cadastre em Configurações › Assinantes.")
        return [completos.get(n.split()[0].casefold(), n) for n in vistos]

    def _capa(self):
        a, p, ed = self.a, self.p, self.ed
        Editor.substituir(a["capa_mun"], r"X{3,}", lista_pt(self.cidades).upper())
        nomes = self._nomes_equipe() or ["Equipe de fiscalização"]
        for n in reversed(nomes[1:]):                               # um parágrafo por analista, no estilo do 1º
            novo = clone(a["an1"]._p)
            a["an1"]._p.addnext(novo)
            Editor.texto_unico(Paragraph(novo, ed.doc._body), n)
            self._pintar_preto(Paragraph(novo, ed.doc._body))
        Editor.texto_unico(a["an1"], nomes[0])
        self._pintar_preto(a["an1"])
        # "Analista 2 MÊS 20XX" -> apenas "MÊS AAAA", com o mesmo recuo dos nomes
        Editor.texto_unico(a["an2"], f"{MESES_UP[p.mes - 1]} {p.ano}")
        self._pintar_preto(a["an2"])
        ind1, ind2 = a["an1"]._p.find(qn("w:pPr") + "/" + qn("w:ind")), a["an2"]._p.find(qn("w:pPr") + "/" + qn("w:ind"))
        if ind1 is not None and ind2 is not None:
            for k in ("left", "right", "firstLine", "hanging"):
                ind2.attrib.pop(qn(f"w:{k}"), None)
            for k, v in ind1.attrib.items():
                ind2.set(k, v)
        # PA-0XX/20XX
        vermelhos = [r for r in a["pa"].runs if r.text == "XX"]
        if len(vermelhos) >= 2:
            if self.op.pa.strip():
                num_pa = int(re.sub(r"\D", "", self.op.pa) or 0)
                vermelhos[0].text = f"{num_pa:02d}"
                Editor.pintar_preto(vermelhos[0])
            vermelhos[1].text = f"{p.ano % 100:02d}"
            Editor.pintar_preto(vermelhos[1])
        if self.op.sei.strip():
            Editor.texto_unico(a["sei"], f"SEI N° {self.op.sei.strip()}")
            self._pintar_preto(a["sei"])

    @staticmethod
    def _pintar_preto(par: Paragraph):
        for r in par.runs:
            Editor.pintar_preto(r)

    # --------------------------------------------------------------------------------------------------
    # Textos fixos com campos variáveis
    # --------------------------------------------------------------------------------------------------
    def _textos_fixos(self):
        a, p = self.a, self.p
        lista = lista_pt(self.cidades)
        sub = Editor.substituir_global
        sub(a["intro2"], r"nos municípios de X+ e X+", f"{'nos municípios' if len(self.cidades) > 1 else 'no município'} de {lista}")
        sub(a["intro2"], r"X{3,} de 20XX", f"{p.nome_mes} de {p.ano}")
        sub(a["meta"], r"2025 é de 372", f"{p.ano} é de {self.an.meta_anual}")
        sub(a["metod"], r"X{6,}", lista)
        sub(a["analise"], r"meta de 372 fiscalizações", f"meta de {self.an.meta_anual} fiscalizações")
        n_mun = self._total_municipios()
        sub(a["para"], r"(?<=entre )26|(?<=dos )26", str(n_mun))

    def _total_municipios(self) -> int:
        metas, _ = self.metas.para_ano(self.p.ano)
        todos = {normalizar(f.cidade) for f in self.ds.fiscalizacoes}
        return max(self.cfg.total_municipios, len(metas), len(todos))

    # --------------------------------------------------------------------------------------------------
    # 4.1  Preparação e planejamento (Tabela 01 + Figura 01)
    # --------------------------------------------------------------------------------------------------
    def _linhas_meta(self, ate_mes: int) -> list[tuple[str, int | None, int]]:
        metas, origem = self.metas.para_ano(self.p.ano)
        cont = Counter(normalizar(f.cidade) for f in self.ds.fiscalizacoes
                       if f.data.year == self.p.ano and f.data.month <= ate_mes)
        nomes = {normalizar(c): c for c in metas}
        for f in self.ds.fiscalizacoes:
            if f.data.year == self.p.ano and f.data.month <= ate_mes:
                nomes.setdefault(normalizar(f.cidade), f.cidade)
        linhas = [(nome.upper(), cont.get(k) or None, next((m for c, m in metas.items() if normalizar(c) == k), 0))
                  for k, nome in nomes.items()]
        linhas.sort(key=lambda x: normalizar(x[0]))
        return linhas

    def _tabela_meta(self, proto_a, proto_b, linhas, titulo_meta: str | None = None) -> object:
        tbl = clone(proto_a)
        trs = tbl.findall(qn("w:tr"))
        cab, corpo = trs[0], trs[1]
        for tr in trs[1:]:
            tbl.remove(tr)
        total = clone(proto_b.findall(qn("w:tr"))[-1])
        for cidade, real, meta in linhas:
            tr = clone(corpo)
            tbl.append(tr)
            c = Table(tbl, self.ed.doc._body).rows[-1].cells
            reescrever_celula_modelo(c[0], cidade)
            reescrever_celula_modelo(c[1], "" if not real else str(real))
            reescrever_celula_modelo(c[2], str(meta))
        tbl.append(total)
        T = Table(tbl, self.ed.doc._body)
        c = T.rows[-1].cells
        reescrever_celula_modelo(c[0], "Total")
        reescrever_celula_modelo(c[1], str(sum(r or 0 for _, r, _ in linhas)))
        reescrever_celula_modelo(c[2], str(sum(m for _, _, m in linhas)))
        if titulo_meta:
            reescrever_celula_modelo(T.rows[0].cells[2], titulo_meta)
        linha_cabecalho(T.rows[0])
        return tbl

    def _secao_41(self):
        a, p, ed = self.a, self.p, self.ed
        anterior = MESES_NOME[(p.mes - 2) % 12]
        self._rotulo_legenda(a["cap01"], "Tabela 01", f"Meta por municípios em {anterior} definidos pela CEEGC")
        linhas = self._linhas_meta(p.mes - 1)
        soma = sum(m for _, _, m in linhas)
        if soma != self.an.meta_anual:
            self.avisos.append(f"A soma da meta por município ({soma}) difere da meta anual de {p.ano} "
                               f"({self.an.meta_anual}). Ajuste 'metas_municipios.csv' (ou a aba Meta_Municipios).")
        origem = self.metas.para_ano(p.ano)[1]
        if origem == "estimada":
            self.avisos.append("Meta por município estimada (não há metas_municipios.csv nem aba Meta_Municipios).")
        ed.remover_entre(a["cap01"]._p, a["fig01"]._p)
        a["cap01"]._p.addnext(self._tabela_meta(self.tp_meta_a, self.tp_meta_b, linhas))
        # Figura 01: imagem/placeholder acima da legenda
        a["fig01"]._p.addprevious(self._figura(1))
        ed.remover_entre(a["fig01"]._p, a["def"]._p)
        sub = Editor.substituir_global
        sub(a["def"], r"X{4,}", lista_pt(self.cidades))
        sub(a["def"], r"no mês de maio", f"no mês de {p.nome_mes}")

    # --------------------------------------------------------------------------------------------------
    # 4.2  Execução (tabelas por dia + relação das NC)
    # --------------------------------------------------------------------------------------------------
    def _rotulo_legenda(self, par: Paragraph, rotulo: str, texto: str):
        runs = [r for r in par.runs if r._r.find(qn("w:drawing")) is None]
        if not runs:
            return
        runs[0].text = f"{rotulo} - {texto}"
        if len(runs) > 1 and runs[0].text.startswith(rotulo):
            # run 0 mantém o negrito do rótulo; run 1 recebe o texto (formatação do modelo)
            runs[0].text = rotulo + " "
            runs[1].text = "- " + texto
            rest = runs[2:]
        else:
            rest = runs[1:]
        for r in rest:
            r.text = ""

    def _tabela_dia(self, dia, fiscs_dia) -> object:
        tbl = clone(self.tp_dia)
        trs = tbl.findall(qn("w:tr"))
        corpo = clone(trs[2])
        for tr in trs[2:]:
            tbl.remove(tr)
        arpe = lista_pt(list(dict.fromkeys(n for f in fiscs_dia for n in f.equipe_arpe)))
        cop = lista_pt(list(dict.fromkeys(f.func_copergas for f in fiscs_dia if f.func_copergas)))
        tc0 = trs[0].findall(qn("w:tc"))[0]
        paras = tc0.findall(qn("w:p"))
        Paragraph(paras[0], None).runs[0].text = f"Equipe Arpe: {arpe}"
        Paragraph(paras[1], None).runs[0].text = f"Funcionário Copergás: {cop}"
        for f in fiscs_dia:
            tr = clone(corpo)
            tbl.append(tr)
            c = Table(tbl, self.ed.doc._body).rows[-1].cells
            reescrever_celula_modelo(c[0], str(f.seq))
            reescrever_celula_modelo(c[1], f.usuario)
            reescrever_celula_modelo(c[2], f.endereco_completo)
            reescrever_celula_modelo(c[3], "CONFORME" if f.conforme else "NÃO CONFORME",
                                     cor=VERDE_OK if f.conforme else VERMELHO_NC)
        T = Table(tbl, self.ed.doc._body)
        for r in T.rows[:2]:
            linha_cabecalho(r)
        for r in T.rows[2:]:
            linha_nao_quebra(r)
        return tbl

    def _tabela_nc(self) -> object:
        tbl = clone(self.tp_nc)
        trs = tbl.findall(qn("w:tr"))
        corpo, total = clone(trs[2]), clone(trs[3])
        for tr in trs[2:]:
            tbl.remove(tr)
        for f, nc in self.ncs:
            tbl.append(clone(corpo))
            c = Table(tbl, self.ed.doc._body).rows[-1].cells
            refs = self.refs_foto.get(nc.num) or []
            txt_foto = "—" if not refs else (f"Foto {refs[0]:02d}" if len(refs) == 1 else
                                              "Fotos " + lista_pt([f"{r:02d}" for r in refs]))
            for cel, txt in zip(c, (str(f.seq), f.usuario, nc.detalhe_limpo, txt_foto,
                                    self.cfg.fundamentacao_artigo, self.cfg.fundamentacao_inciso, "1")):
                reescrever_celula_modelo(cel, txt)
            linha_nao_quebra(Table(tbl, self.ed.doc._body).rows[-1])
        tbl.append(total)
        cells = Table(tbl, self.ed.doc._body).rows[-1].cells
        reescrever_celula_modelo(cells[-1], str(len(self.ncs)))
        for r in Table(tbl, self.ed.doc._body).rows[:2]:
            linha_cabecalho(r)
        return tbl

    def _secao_42(self):
        a, p, ed, nm = self.a, self.p, self.ed, self.nm
        dias = sorted({f.data for f in self.fiscs})
        nm_dias = nm.refs(nm.dias, artigo=True)
        verbo = "apresenta" if len(dias) == 1 else "apresentam"
        Editor.substituir_global(a["exec"], r"As Tabelas 02 e 03 apresentam os clientes fiscalizados em cada um dos dias que foram "
                                 r"realizadas fiscalizações no mês",
                                 f"{nm_dias} {verbo} os clientes fiscalizados em cada um dos dias que foram realizadas "
                                 f"fiscalizações no mês")
        Editor.substituir_global(a["exec"], r"relacionadas na Tabela 03", f"relacionadas na Tabela {nm.t(nm.nc)}")

        # --- tabelas por dia
        novos = []
        for n, dia in zip(nm.dias, dias):
            do_dia = [f for f in self.fiscs if f.data == dia]
            novos.append(ed.legenda(self.m_cap, f"Tabela {nm.t(n)}", f"Fiscalizações realizadas no dia {_data_br(dia)}")._p)
            novos.append(self._tabela_dia(dia, do_dia))
            novos.append(self._espaco())
        ed.remover_entre(a["cap_dia1"]._p, a["cap_nc"]._p)
        a["cap_dia1"]._p.getparent().remove(a["cap_dia1"]._p)
        for el in novos:
            a["cap_nc"]._p.addprevious(el)

        # --- relação das NC
        self._rotulo_legenda(a["cap_nc"], f"Tabela {nm.t(nm.nc)}", "Relação das Não Conformidades constatadas")
        if self.ncs:
            ed.remover_entre(a["cap_nc"]._p, a["nc_leg"]._p)
            a["cap_nc"]._p.addnext(self._tabela_nc())
            Editor.substituir_global(a["nc_leg"], r"Tabela 06", f"Tabela {nm.t(nm.nc)}")
        else:
            ed.remover_entre(a["cap_nc"]._p, a["h_monit"]._p)
            a["cap_nc"]._p.getparent().remove(a["cap_nc"]._p)
            novo = ed.paragrafo_texto(self.m_txt, "Não foram constatadas não conformidades nas fiscalizações realizadas "
                                                   f"em {p.nome_mes} de {p.ano}.")
            a["h_monit"]._p.addprevious(novo._p)

    # --------------------------------------------------------------------------------------------------
    # 4.3 e 5
    # --------------------------------------------------------------------------------------------------
    def _secao_43_e_5(self):
        a, p, ed, nm = self.a, self.p, self.ed, self.nm
        sub = Editor.substituir_global
        sub(a["p43"], r"A Tabela 04", f"A Tabela {nm.t(nm.meta_atual)}")
        sub(a["p43"], r"no mês de junho", f"no mês de {p.nome_mes}")
        self._rotulo_legenda(a["cap_meta2"], f"Tabela {nm.t(nm.meta_atual)}",
                             f"Meta por municípios em {p.nome_mes} definidos pela CEEGC")
        linhas = self._linhas_meta(p.mes)
        soma = sum(m for _, _, m in linhas)
        fim = ed.achar(inicio="DETERMINAÇÕES GERAIS", apos=a["cap_meta2"]._p)
        ed.remover_entre(a["cap_meta2"]._p, fim._p)
        a["cap_meta2"]._p.addnext(self._tabela_meta(self.tp_meta2a, self.tp_meta2b, linhas,
                                                    titulo_meta=f"Amostra de fiscalizações Meta {soma}"))
        # 5. determinações
        n = len(self.ncs)
        if n == 0:
            Editor.texto_unico(a["deter"], "Não foram constatadas não conformidades nas fiscalizações realizadas no período. "
                                           "Reiteramos a importância do monitoramento contínuo e da manutenção dos sistemas "
                                           "de distribuição de gás natural para a eficiência dos serviços prestados.")
        elif n > 1:
            sub(a["deter"], r"quanto à (?=não conformidade)", "quanto às ")
            sub(a["deter"], r"não conformidade(?= mencionada)", "não conformidades")
            sub(a["deter"], r"mencionada(?=,)", "mencionadas")

    # --------------------------------------------------------------------------------------------------
    # Apêndice 1
    # --------------------------------------------------------------------------------------------------
    @staticmethod
    def _tam_foto(path: Path, larg_max=8.4, alt_max=6.0) -> tuple[float, float]:
        with Image.open(path) as im:
            w, h = im.size
        esc = min(larg_max / w, alt_max / h)
        return w * esc, h * esc

    def _tabelas_fotos(self, itens: list[FotoItem], inicio: int = 1) -> list:
        ed, cfg = self.ed, self.cfg
        cols, por_tab = cfg.fotos_colunas, cfg.fotos_por_tabela
        out = []
        for i in range(0, len(itens), por_tab):
            bloco = itens[i:i + por_tab]
            if i:
                quebra = ed.doc.add_page_break()
                out.append(quebra._p)
            t = ed.doc.add_table(rows=-(-len(bloco) // cols), cols=cols)
            bordas_tabela(t, "000000", 6)
            margens_tabela(t, 0.12)
            layout_fixo(t, [9.0] * cols)
            for k, it in enumerate(bloco):
                cel = t.cell(k // cols, k % cols)
                par = cel.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(4)
                w, h = self._tam_foto(it.path)
                par.add_run().add_picture(str(it.path), width=Cm(w), height=Cm(h))
                cap = cel.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(4)
                r = cap.add_run(f"Foto {inicio + i + k}: {it.legenda}")
                r.font.size, r.font.name = Pt(8.5), "Calibri"
            for row in t.rows:
                linha_nao_quebra(row)
            out.append(t._tbl)
        return out

    def _apendice1(self):
        a, ed = self.a, self.ed
        ed.remover_entre(a["ap1"]._p, a["ap2"]._p)
        novos = []
        quebra = ed.doc.add_page_break()
        a["ap1"]._p.addprevious(quebra._p)
        if self.itens_fotos:
            novos += self._tabelas_fotos(self.itens_fotos)
        else:
            novos.append(ed.paragrafo_texto(self.m_txt, "Não há registros fotográficos de não conformidades neste período.")._p)
        if self.op.incluir_condicoes_gerais:
            gerais = []
            for cidade, vid in self.vids.items():
                gerais += [FotoItem(c, f"Condições gerais – {cidade}") for c in self.fotos.fotos_gerais(vid, cidade)]
            if gerais:
                novos.append(ed.doc.add_page_break()._p)
                h = ed.paragrafo_texto(self.m_h, "APÊNDICE 1-A - REGISTROS DAS CONDIÇÕES GERAIS")
                novos.append(h._p)
                novos += self._tabelas_fotos(gerais)
        novos.append(ed.doc.add_page_break()._p)
        for el in novos:
            a["ap2"]._p.addprevious(el)

    # --------------------------------------------------------------------------------------------------
    # Apêndice 2
    # --------------------------------------------------------------------------------------------------
    def _espaco(self, pt: float = 6):
        p = self.ed.doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        sp = OxmlElement("w:spacing")
        for k, v in (("before", "0"), ("after", "0"), ("line", str(int(pt * 20))), ("lineRule", "exact")):
            sp.set(qn(f"w:{k}"), v)
        ppr.append(sp)
        return p._p

    def _figura(self, n: int):
        """Imagem do mapa (assets/mapas/figuraNN.*) ou área reservada em vermelho."""
        ed = self.ed
        arq = self._arquivo_mapa(n)
        if arq:
            return ed.paragrafo_imagem(str(arq), 15.5, altura_max_cm=11.5)._p
        self.avisos.append(f"Figura {n:02d} (mapa) não encontrada em assets/mapas/ — foi deixada uma área reservada no relatório.")
        p = ed.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ppr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        for lado in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{lado}")
            for k, v in (("val", "dashed"), ("sz", "6"), ("space", "8"), ("color", "FF0000")):
                e.set(qn(f"w:{k}"), v)
            bd.append(e)
        ppr.append(bd)
        p.paragraph_format.space_before = Pt(50)
        p.paragraph_format.space_after = Pt(50)
        r = p.add_run(f"[Inserir aqui o mapa da Figura {n:02d}. Salve o arquivo como figura{n:02d}.png em assets/mapas/ "
                      f"para inserção automática.]")
        r.font.size, r.font.color.rgb, r.italic = Pt(9), RGBColor(0xFF, 0, 0), True
        return p._p

    def _arquivo_mapa(self, n: int) -> Path | None:
        d = self.cfg.mapas_dir
        if not d.exists():
            return None
        alvo = {f"figura{n:02d}", f"figura_{n:02d}", f"figura {n:02d}", f"fig{n:02d}", f"figura{n}", f"figura_{n}"}
        for arq in sorted(d.iterdir()):
            if arq.suffix.lower() in (".png", ".jpg", ".jpeg") and arq.stem.casefold() in alvo:
                return arq
        return None

    # ---------- tabelas no estilo da planilha (faixa azul + zebra) ----------
    def _banda(self, numero: int | None, texto: str, larguras: list[float]):
        t = self.ed.doc.add_table(rows=1, cols=2)
        bordas_tabela(t, "808080", 4)
        margens_tabela(t, 0.05)
        total = sum(larguras)
        layout_fixo(t, [0.8, total - 0.8])
        for c, txt, al in ((t.cell(0, 0), "" if numero is None else str(numero), "center"),
                           (t.cell(0, 1), texto, "center")):
            sombrear(c, AZUL_BANDA)
            escrever(c, txt, negrito=True, tam=7, alinh=al)
        return t._tbl

    def _tab_excel(self, cabecalho: list[str], linhas: list[list], larguras: list[float], total_ultima=False,
                   zebra=True, alinhar_esq=(0,), tam=7, negrito_cols=()):
        t = self.ed.doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
        bordas_tabela(t, "808080", 4)
        margens_tabela(t, 0.04)
        layout_fixo(t, larguras)
        for j, h in enumerate(cabecalho):
            escrever(t.cell(0, j), h, negrito=True, tam=tam)
        linha_cabecalho(t.rows[0])
        for i, lin in enumerate(linhas, start=1):
            ultima = total_ultima and i == len(linhas)
            for j, v in enumerate(lin):
                c = t.cell(i, j)
                escrever(c, "" if v is None else v, negrito=(ultima or j in negrito_cols), tam=tam,
                         alinh="left" if j in alinhar_esq else "center")
                if zebra and i % 2 == 0 and not ultima:
                    sombrear(c, CINZA_ZEBRA)
            linha_nao_quebra(t.rows[i])
        return t._tbl

    @staticmethod
    def _fmt(v, casas=0, branco_zero=False) -> str:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return ""
        if branco_zero and not v:
            return ""
        return f"{v:.{casas}f}".replace(".", ",")

    def _tabela_nicho(self) -> list:
        an, n = self.an, self.p.mes
        k = an.resumo_meta()
        larg_res = [2.1, 1.3, 2.1, 1.3, 2.4, 1.6, 2.4, 1.3]
        res = self.ed.doc.add_table(rows=1, cols=8)
        bordas_tabela(res, "808080", 4)
        layout_fixo(res, larg_res)
        vals = ["META ANUAL:", str(k["meta_anual"]), "REALIZADO:", str(k["realizado"]), "% REAL/META:",
                self._fmt(k["pct"], 2), "META MENSAL:", self._fmt(k["meta_mensal"], 0)]
        for j, v in enumerate(vals):
            c = res.cell(0, j)
            if j % 2 == 0:
                sombrear(c, AZUL_BANDA)
            escrever(c, v, negrito=True, tam=7, alinh="left" if j % 2 == 0 else "center")
        w = [3.0] + [0.95] * 12 + [1.2, 1.2]
        seg = an.por_segmento_mes()
        nomes = [s for s in ORDEM_SEGMENTOS] + [s for s in seg.index if s not in ORDEM_SEGMENTOS and s != "Total"]
        linhas = []
        for s in nomes:
            r = seg.loc[s] if s in seg.index else None
            meses = [self._fmt(int(r[m]) if r is not None else 0, 0, True) if i < n else "" for i, m in enumerate(MESES_ABREV)]
            tot = int(r["TOTAL"]) if r is not None else 0
            linhas.append([s] + meses + [str(tot), self._fmt(float(r["%"]) if r is not None else 0.0, 2)])
        tr = seg.loc["Total"]
        linhas.append(["Total"] + [str(int(tr[m])) if i < n else "0" for i, m in enumerate(MESES_ABREV)]
                      + [str(int(tr["TOTAL"])), "100,00"])
        return [res._tbl, self._banda(1, "QUANTITATIVO DE FISCALIZAÇÕES REALIZADAS MENSALMENTE:", w),
                self._tab_excel(["Setor"] + MESES_ABREV + ["TOTAL", "%"], linhas, w, total_ultima=True)]

    def _tabela_municipio(self) -> list:
        an, n = self.an, self.p.mes
        cid = an.por_cidade_mes()
        metas, _ = self.metas.para_ano(self.p.ano)
        nomes = {normalizar(c): titulo_pt(c) for c in metas}
        for c in cid.index:
            if c != "Total":
                nomes[normalizar(c)] = c
        por_norm = {normalizar(c): c for c in cid.index if c != "Total"}
        linhas = []
        for k in sorted(nomes, key=lambda x: x):
            r = cid.loc[por_norm[k]] if k in por_norm else None
            meses = [self._fmt(int(r[m]) if r is not None else 0, 0, True) if i < n else "" for i, m in enumerate(MESES_ABREV)]
            linhas.append([nomes[k]] + meses + [str(int(r["TOTAL"]) if r is not None else 0)])
        tr = cid.loc["Total"]
        linhas.append(["Total"] + [str(int(tr[m])) if i < n else "0" for i, m in enumerate(MESES_ABREV)] + [str(int(tr["TOTAL"]))])
        w = [3.6] + [0.95] * 12 + [1.3]
        return [self._banda(2, f"ACOMPANHAMENTO MENSAL DO QUANTITATIVO DE CLIENTES POR MUNICÍPIO FISCALIZADOS EM {self.p.ano}:", w),
                self._tab_excel(["Município"] + MESES_ABREV + ["TOTAL"], linhas, w, total_ultima=True)]

    def _tabela_pct_meta(self) -> list:
        ev, n = self.an.evolucao_mensal(), self.p.mes
        linha = ["% FISC/META"] + [self._fmt(ev.pct_meta.iloc[i], 1) if i < n else "0,0" for i in range(12)] \
            + [self._fmt(float(ev.pct_meta.iloc[:n].sum()), 1)]
        w = [2.6] + [1.0] * 12 + [1.4]
        return [self._banda(5, "PERCENTUAL DAS FISCALIZAÇÕES EM RELAÇÃO A META E AO REALIZADO:", w),
                self._tab_excel([""] + MESES_ABREV + ["TOTAL"], [linha], w, zebra=False)]

    def _tabela_dias(self) -> list:
        dm = self.an.dias_e_media()
        linhas = []
        for rot in ("Dias programados", "Dias usados", "Fiscs / dia"):
            if rot in dm.index:
                casas = 2 if rot == "Fiscs / dia" else 0
                linhas.append([rot] + [self._fmt(dm.loc[rot, m], casas) for m in MESES_ABREV]
                              + [self._fmt(dm.loc[rot, "TOTAL"], casas)])
        w = [2.6] + [1.0] * 12 + [1.4]
        return [self._banda(6, "DIAS UTILIZADOS NA FISCALIZAÇÃO E MÉDIA DE FISCALIZAÇÃO POR DIA:", w),
                self._tab_excel([""] + MESES_ABREV + ["TOTAL"], linhas, w)]

    def _tabela_acomp(self, df: pd.DataFrame, titulo: str, numero: int) -> list:
        linhas = []
        for i, (rot, r) in enumerate(df.iterrows(), start=1):
            linhas.append([str(i), rot] + [self._fmt(r[m], 0) for m in MESES_ABREV] + [self._fmt(r["TOTAL"], 0)])
        w = [0.6, 4.2] + [0.95] * 12 + [1.3]
        return [self._banda(numero, titulo, w),
                self._tab_excel(["", ""] + MESES_ABREV + ["TOTAL"], linhas, w, alinhar_esq=(1,))]

    def _tabela_frequencia(self) -> list:
        fr = self.an.frequencia_nc(top=4, outros=True)
        det = int(self.an.frequencia_nc()["Frequência"].sum()) if len(self.an.nc) else 0
        linhas = [[r["Não conformidade"], str(int(r["Frequência"])), self._fmt(float(r["%"]), 2)] for _, r in fr.iterrows()]
        linhas.append(["Total", str(det), "100,00" if det else "0,00"])
        w = [10.0, 3.0, 3.0]
        return [self._banda(10, f"NÃO CONFORMIDADES MAIS FREQUENTES EM {self.p.ano}:", w),
                self._tab_excel(["NÃO CONFORMIDADE", "FREQUÊNCIA", "%"], linhas, w, total_ultima=True)]

    def _tabela_mun_mes(self) -> object:
        pend = set(self.an.cidades_com_pendencia())
        cont = Counter(f.cidade for f, _ in self.ncs)
        t = self.ed.doc.add_table(rows=2 + len(self.cidades) + 1, cols=3)
        bordas_tabela(t, "000000", 6)
        layout_fixo(t, [6.0, 2.0, 6.0])
        top = t.cell(0, 0).merge(t.cell(0, 2))
        escrever(top, "MUNICÍPIOS FISCALIZADOS NO MÊS", negrito=True, tam=9, fonte="Calibri")
        for j, h in enumerate(("MUNICÍPIOS", "COR", "QUANTITATIVO DE NÃO CONFORMIDADES")):
            sombrear(t.cell(1, j), "000000")
            escrever(t.cell(1, j), h, tam=8, cor="FFFFFF", fonte="Calibri")
        for i, cid in enumerate(self.cidades, start=2):
            escrever(t.cell(i, 0), cid, tam=9, alinh="left", fonte="Calibri")
            sombrear(t.cell(i, 1), "F4B183" if cid in pend else "70AD47")
            escrever(t.cell(i, 1), "", tam=9)
            escrever(t.cell(i, 2), str(cont.get(cid, 0)), tam=9, alinh="right", fonte="Calibri")
        ult = len(self.cidades) + 2
        escrever(t.cell(ult, 0), "Total", negrito=True, tam=9, alinh="left", fonte="Calibri")
        escrever(t.cell(ult, 2), str(len(self.ncs)), negrito=True, tam=9, alinh="right", fonte="Calibri")
        return t._tbl

    def _grafico_com_tabela(self, titulo_banda: str, numero: int, fig, linhas: list[list], cores: list[str | None]) -> list:
        w = [2.6] + [1.0] * 12
        img = self.ed.paragrafo_imagem(ChartFactory.para_png(fig), 16.0, manter_com_proximo=True)._p
        tab = self.ed.doc.add_table(rows=len(linhas), cols=13)
        bordas_tabela(tab, "BFBFBF", 4)
        margens_tabela(tab, 0.03)
        layout_fixo(tab, w)
        for i, lin in enumerate(linhas):
            for j, v in enumerate(lin):
                c = tab.cell(i, j)
                escrever(c, v, tam=7, alinh="left" if j == 0 else "center", negrito=(j == 0))
                if j == 0 and cores[i]:
                    escrever(c, "■ " + v, tam=7, alinh="left", negrito=True, cor=cores[i])
        return [self._banda(numero, titulo_banda, w), img, tab._tbl]

    def _apendice2(self):
        a, p, ed, nm, an = self.a, self.p, self.ed, self.nm, self.an
        cf = ChartFactory(an)
        ev, n = an.evolucao_mensal(), p.mes
        novos: list = []

        def txt(t):
            novos.append(ed.paragrafo_texto(self.m_txt, t)._p)

        def cap(rot, t):
            novos.append(ed.legenda(self.m_cap, rot, t)._p)

        def add(*els):
            for e in els:
                novos.append(e)
            novos.append(self._espaco())

        txt(self.narr.nicho_municipio())
        cap(f"Tabela {nm.t(nm.segmento)}", "Fiscalizações por nicho e mês")
        add(*self._tabela_nicho())
        cap(f"Tabela {nm.t(nm.municipio)}", "Fiscalizações por município e mês")
        add(*self._tabela_municipio())
        for t_ in self.narr.evolucao():
            txt(t_)
        meses = lambda s: [self._fmt(s.iloc[i], 0) if i < n else "" for i in range(12)]  # noqa: E731
        cap("Gráfico 01", "Quantitativo de fiscalizações por mês")
        add(*self._grafico_com_tabela("GRÁFICO DA EVOLUÇÃO DAS FISCALIZAÇÕES MENSAIS META X REALIZADO:", 3, cf.mensal(),
                                      [["Mês"] + MESES_ABREV, ["Realizado"] + meses(ev.realizado),
                                       ["Meta"] + [self._fmt(ev.meta.iloc[i], 0) for i in range(12)]],
                                      [None, "4472C4", "ED7D31"]))
        cap("Gráfico 02", "Evolução das fiscalizações no ano")
        add(*self._grafico_com_tabela("GRÁFICO DA EVOLUÇÃO DAS FISCALIZAÇÕES MENSAIS ACUMULADAS: META X REALIZADO:", 4, cf.acumulado(),
                                      [["Mês"] + MESES_ABREV, ["Realizado"] + meses(ev.acum_real),
                                       ["Meta"] + [self._fmt(ev.acum_meta.iloc[i], 0) for i in range(12)]],
                                      [None, "4472C4", "ED7D31"]))
        cap(f"Tabela {nm.t(nm.pct_meta)}", "Percentual de fiscalizações baseado na meta")
        add(*self._tabela_pct_meta())
        cap(f"Tabela {nm.t(nm.media_dia)}", "Média do quantitativo de fiscalizações por dia")
        add(*self._tabela_dias())
        txt(self.narr.nao_conformidades())
        cap(f"Tabela {nm.t(nm.clientes)}", "Acompanhamento clientes não conformes")
        add(*self._tabela_acomp(an.acompanhamento_clientes(), "ACOMPANHAMENTO MENSAL DOS CLIENTES NÃO CONFORMES:", 7))
        cap(f"Tabela {nm.t(nm.nc_acomp)}", "Acompanhamento não conformidades")
        add(*self._tabela_acomp(an.acompanhamento_nc(), "ACOMPANHAMENTO MENSAL DAS NÃO CONFORMIDADES:", 8))
        ac = an.acompanhamento_nc()
        lin = lambda rot: [self._fmt(ac.loc[rot, m], 0) for m in MESES_ABREV]  # noqa: E731
        cap("Gráfico 03", "Acompanhamento mensal não conformidades")
        add(*self._grafico_com_tabela(f"ACOMPANHAMENTO MENSAL DAS NÃO CONFORMIDADES:", 9, cf.nc_evolucao(),
                                      [["Mês"] + MESES_ABREV, ["NC detectadas"] + lin("NC detectadas no mês"),
                                       ["NC sanadas"] + lin("NC sanadas"), ["NC pendentes"] + lin("Saldo NC pendentes")],
                                      [None, "ED7D31", "70AD47", "FF0000"]))
        cap(f"Tabela {nm.t(nm.frequencia)}", "Frequência de não conformidades")
        add(*self._tabela_frequencia())
        txt(self.narr.mapas())
        for fig_n, titulo in ((2, "Mapa Pernambuco situação das fiscalizações e não conformidades"),
                              (3, "Mapa gasoduto principal situação das fiscalizações e não conformidades"),
                              (4, "Mapa municípios fiscalizados no mês situação das não conformidades")):
            novos.append(self._figura(fig_n))
            cap(f"Figura {fig_n:02d}", titulo)
            if fig_n == 4:
                novos.append(self._tabela_mun_mes())
                novos.append(self._espaco(10))

        ed.remover_entre(a["ap2_intro"]._p, a["assin"]._p)
        for el in novos:
            a["assin"]._p.addprevious(el)
